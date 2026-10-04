"""
Per-file plan limits: size, spreadsheet rows and PDF pages, checked before any heavy work starts.

Counting is cheap (PDF page tables and spreadsheet dimensions, not full parsing), so a file that is
too big for the plan is turned away in milliseconds with a message saying what to do.
"""
from __future__ import annotations

import io
import zipfile
from typing import Dict, Optional

from fastapi import HTTPException

UNLIMITED = -1


def _fmt(n: int) -> str:
    return f"{n:,}"


def is_pdf(data: bytes, filename: str = "") -> bool:
    return data[:5] == b"%PDF-" or (filename or "").lower().endswith(".pdf")


def count_pdf_pages(data: bytes) -> Optional[int]:
    try:
        import fitz  # PyMuPDF

        with fitz.open(stream=data, filetype="pdf") as doc:
            return int(doc.page_count)
    except Exception:
        pass
    try:
        from PyPDF2 import PdfReader

        return len(PdfReader(io.BytesIO(data), strict=False).pages)
    except Exception:
        return None


def count_spreadsheet_rows(data: bytes, filename: str = "", stop_after: Optional[int] = None) -> Optional[int]:
    """Rows across all sheets (or lines in a CSV). None when the file isn't a spreadsheet we can read."""
    name = (filename or "").lower()
    if name.endswith((".csv", ".tsv", ".txt")):
        n = data.count(b"\n")
        if data and not data.endswith(b"\n"):
            n += 1
        return n
    if data[:2] == b"PK" and (name.endswith((".xlsx", ".xlsm")) or not name):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if not any(n.startswith("xl/worksheets/") for n in zf.namelist()):
                    return None
        except zipfile.BadZipFile:
            return None
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                total = 0
                for ws in wb.worksheets:
                    rows = ws.max_row  # from the sheet's stored dimension; cheap
                    if rows is None:
                        rows = 0
                        for _ in ws.iter_rows(values_only=True):
                            rows += 1
                            if stop_after is not None and total + rows > stop_after:
                                break
                    total += int(rows or 0)
                    if stop_after is not None and total > stop_after:
                        return total
                return total
            finally:
                wb.close()
        except Exception:
            return None
    if name.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        try:
            import xlrd

            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                return sum(book.sheet_by_index(i).nrows for i in range(book.nsheets))
            finally:
                book.release_resources()
        except Exception:
            return None
    return None


def check_file(limits: Dict[str, int], data: bytes, filename: str = "", plan_name: str = "your plan") -> None:
    """Raise 413 when one uploaded file is over the plan's size, row or page limit."""
    size_mb_limit = limits.get("file_size_mb")
    if size_mb_limit is not None and size_mb_limit >= 0 and len(data) > size_mb_limit * 1024 * 1024:
        raise HTTPException(
            status_code=413,
            detail=f"File size ({len(data) / (1024 * 1024):.1f}MB) exceeds {size_mb_limit}MB limit on {plan_name}.",
        )

    if is_pdf(data, filename):
        page_limit = limits.get("pdf_pages")
        if page_limit is not None and page_limit >= 0:
            pages = count_pdf_pages(data)
            if pages is not None and pages > page_limit:
                raise HTTPException(
                    status_code=413,
                    detail=(
                        f"This PDF has {_fmt(pages)} pages; {plan_name} allows {_fmt(page_limit)} per file. "
                        "Split it into smaller PDFs or upgrade your plan."
                    ),
                )
        return

    row_limit = limits.get("spreadsheet_rows")
    if row_limit is not None and row_limit >= 0:
        rows = count_spreadsheet_rows(data, filename, stop_after=row_limit)
        if rows is not None and rows > row_limit:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"This spreadsheet has more than {_fmt(row_limit)} rows, the most {plan_name} allows per file. "
                    "Split it or upgrade your plan."
                ),
            )


def ocr_page_cap(limits: Dict[str, int], requested: Optional[int], server_default: int) -> int:
    """Pages OCR may read in one file: the smallest of what was asked, the server default and the plan."""
    cap = limits.get("ocr_pages", UNLIMITED)
    values = [v for v in (requested, server_default) if v is not None and v > 0]
    if cap is not None and cap >= 0:
        values.append(cap)
    return max(1, min(values)) if values else 1


__all__ = ["check_file", "count_pdf_pages", "count_spreadsheet_rows", "ocr_page_cap", "is_pdf"]
