import io
import logging
import re
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)


@dataclass
class IngestLimits:
    max_bytes: int = 10 * 1024 * 1024
    max_chars: int = 60_000
    max_pages: int = 25
    max_sheets: int = 5
    max_rows_per_sheet: int = 250
    max_cells_per_sheet: int = 10_000


def _normalize_whitespace(s: str) -> str:
    if not s:
        return ""
    s = s.replace("\x00", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _clip(s: str, max_chars: int) -> str:
    if s is None:
        return ""
    if len(s) <= max_chars:
        return s
    return s[: max_chars - 200] + "\n\n...[truncated]..."


class IngestionService:
    def __init__(self, limits: Optional[IngestLimits] = None):
        self.limits = limits or IngestLimits()

    def ingest(self, filename: str, content: bytes) -> Dict[str, Any]:
        if not filename:
            raise ValueError("filename is required")
        if content is None:
            raise ValueError("content is required")
        if len(content) > self.limits.max_bytes:
            raise ValueError(f"File too large (>{self.limits.max_bytes} bytes)")

        ext = filename.lower().split(".")[-1] if "." in filename else ""
        if ext == "md":
            return self._ingest_markdown(filename, content)
        if ext == "docx":
            return self._ingest_docx(filename, content)
        if ext in ("xlsx", "xls"):
            return self._ingest_excel(filename, content, ext)
        if ext == "pptx":
            return self._ingest_pptx(filename, content)
        if ext == "pdf":
            return self._ingest_pdf(filename, content)

        raise ValueError(f"Unsupported file type: .{ext}")

    def _ingest_markdown(self, filename: str, content: bytes) -> Dict[str, Any]:
        text = content.decode("utf-8", errors="replace")
        text = _clip(text, self.limits.max_chars)
        return {
            "filename": filename,
            "type": "md",
            "extracted_text": text,
            "meta": {"chars": len(text)},
        }

    def _ingest_docx(self, filename: str, content: bytes) -> Dict[str, Any]:
        try:
            from docx import Document
        except Exception as e:
            raise ValueError("python-docx is not available") from e

        doc = Document(io.BytesIO(content))
        parts: List[str] = []

        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                parts.append(t)

        for table in doc.tables:
            for row in table.rows:
                row_vals = []
                for cell in row.cells:
                    t = _normalize_whitespace(cell.text or "")
                    if t:
                        row_vals.append(t)
                if row_vals:
                    parts.append(" | ".join(row_vals))

        text = "\n".join(parts).strip()
        text = _clip(text, self.limits.max_chars)

        return {
            "filename": filename,
            "type": "docx",
            "extracted_text": text,
            "meta": {
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "chars": len(text),
            },
        }

    def _ingest_excel(self, filename: str, content: bytes, ext: str) -> Dict[str, Any]:
        try:
            import pandas as pd
        except Exception as e:
            raise ValueError("pandas is not available") from e

        sheets: Dict[str, Any] = {}
        extracted_blocks: List[str] = []

        if ext == "xlsx":
            engine = "openpyxl"
        else:
            engine = None

        xl = pd.ExcelFile(io.BytesIO(content), engine=engine)
        sheet_names = xl.sheet_names[: self.limits.max_sheets]

        for sn in sheet_names:
            df = xl.parse(sn, nrows=self.limits.max_rows_per_sheet)
            if df is None:
                continue

            # Keep a compact preview
            preview_rows = min(len(df), 10)
            preview = df.head(preview_rows).fillna("")
            preview_csv = preview.to_csv(index=False)

            sheets[sn] = {
                "rows": int(len(df)),
                "cols": int(len(df.columns)),
                "preview_csv": preview_csv,
            }

            extracted_blocks.append(f"# Sheet: {sn}\nRows: {len(df)} Cols: {len(df.columns)}\n\n{preview_csv}")

        text = "\n\n".join(extracted_blocks).strip()
        text = _clip(text, self.limits.max_chars)

        return {
            "filename": filename,
            "type": "excel",
            "extracted_text": text,
            "meta": {
                "sheets": sheets,
                "sheet_count": len(sheet_names),
                "chars": len(text),
            },
        }

    def _ingest_pptx(self, filename: str, content: bytes) -> Dict[str, Any]:
        try:
            from pptx import Presentation
        except Exception as e:
            raise ValueError("python-pptx is not available") from e

        prs = Presentation(io.BytesIO(content))
        slide_texts: List[str] = []

        for idx, slide in enumerate(prs.slides[:100], start=1):
            parts: List[str] = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                t = _normalize_whitespace(shape.text or "")
                if t:
                    parts.append(t)
            if parts:
                slide_texts.append(f"Slide {idx}: " + " | ".join(parts))

        text = "\n".join(slide_texts).strip()
        text = _clip(text, self.limits.max_chars)

        return {
            "filename": filename,
            "type": "pptx",
            "extracted_text": text,
            "meta": {
                "slides": len(prs.slides),
                "chars": len(text),
            },
        }

    def _ingest_pdf(self, filename: str, content: bytes) -> Dict[str, Any]:
        text = ""
        meta: Dict[str, Any] = {}

        # Prefer pdfplumber for layout-ish extraction; fallback to PyMuPDF.
        try:
            import pdfplumber

            pages_text: List[str] = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                meta["pages"] = len(pdf.pages)
                for p in pdf.pages[: self.limits.max_pages]:
                    t = p.extract_text() or ""
                    t = t.strip()
                    if t:
                        pages_text.append(t)
            text = "\n\n".join(pages_text).strip()
        except Exception as e:
            logger.info(f"pdfplumber extraction failed, falling back to PyMuPDF: {e}")

        if not text:
            try:
                import fitz  # PyMuPDF

                doc = fitz.open(stream=content, filetype="pdf")
                meta["pages"] = doc.page_count
                pages_text = []
                for i in range(min(doc.page_count, self.limits.max_pages)):
                    pages_text.append(doc.load_page(i).get_text("text").strip())
                text = "\n\n".join([t for t in pages_text if t]).strip()
            except Exception as e:
                raise ValueError("Failed to extract PDF text") from e

        text = _clip(text, self.limits.max_chars)
        meta["chars"] = len(text)

        return {
            "filename": filename,
            "type": "pdf",
            "extracted_text": text,
            "meta": meta,
        }


def build_ingestion_prompt_block(ingested: Dict[str, Any]) -> str:
    if not ingested:
        return ""
    filename = ingested.get("filename") or "uploaded_file"
    ftype = ingested.get("type") or "file"
    meta = ingested.get("meta") or {}
    extracted = ingested.get("extracted_text") or ""

    header = f"UPLOADED FILE CONTEXT\n- filename: {filename}\n- type: {ftype}\n- meta: {meta}\n\n"
    return header + extracted
