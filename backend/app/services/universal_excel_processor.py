import io
import re
from collections import defaultdict
from datetime import datetime
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl


class UniversalExcelProcessor:
    def __init__(self, file_bytes: bytes, filename: str):
        self.file_bytes = file_bytes
        self.filename = filename
        self.file_size_mb = len(file_bytes) / (1024 * 1024)

        stream = io.BytesIO(file_bytes)
        self.wb_values = openpyxl.load_workbook(stream, read_only=True, data_only=True)

        stream.seek(0)
        self.wb_formulas = openpyxl.load_workbook(stream, read_only=True, data_only=False)

    def close(self) -> None:
        try:
            self.wb_values.close()
        except Exception:
            pass
        try:
            self.wb_formulas.close()
        except Exception:
            pass

    def _iter_sample_row_ranges(self, max_row: int) -> List[Tuple[int, int]]:
        if max_row <= 0:
            return []
        return [
            (1, min(100, max_row)),
            (max(1, max_row // 2 - 50), min(max_row, max_row // 2 + 50)),
            (max(1, max_row - 100), max_row),
        ]

    def diagnose_sheet_formulas(self, sheet_name: str) -> Dict[str, Any]:
        ws_f = self.wb_formulas[sheet_name]
        ws_v = self.wb_values[sheet_name]

        diag: Dict[str, Any] = {
            "sheet_name": sheet_name,
            "has_formulas": False,
            "formula_count": 0,
            "missing_cached_count": 0,
            "risk_level": "ok",
            "reason": None,
            "confidence": 1.0,
        }

        max_row = int(getattr(ws_f, "max_row", 0) or 0)
        max_col = int(getattr(ws_f, "max_column", 0) or 0)
        if max_row <= 0 or max_col <= 0:
            return diag

        sample_ranges = self._iter_sample_row_ranges(max_row)

        formula_coords: List[str] = []
        none_count = 0
        empty_str_count = 0
        numeric_formula_count = 0
        any_numeric_value = False

        for start_r, end_r in sample_ranges:
            for row in ws_f.iter_rows(min_row=start_r, max_row=end_r, min_col=1, max_col=min(max_col, 80), values_only=False):
                for cell in row:
                    v = cell.value
                    is_formula = getattr(cell, "data_type", None) == "f" or (isinstance(v, str) and v.startswith("="))
                    if not is_formula:
                        continue

                    diag["has_formulas"] = True
                    formula_coords.append(cell.coordinate)

                    ftxt = str(v) if isinstance(v, str) else ""
                    up = ftxt.upper()
                    if any(t in up for t in ("SUM(", "AVERAGE(", "MIN(", "MAX(", "COUNT(", "XLOOKUP(", "VLOOKUP(")) or ("+" in ftxt) or ("*" in ftxt) or ("/" in ftxt) or ("-" in ftxt[1:]):
                        numeric_formula_count += 1

                    try:
                        cv = ws_v[cell.coordinate].value
                    except Exception:
                        cv = None

                    if cv is None:
                        none_count += 1
                    elif cv == "":
                        empty_str_count += 1
                    elif isinstance(cv, (int, float)):
                        any_numeric_value = True

        diag["formula_count"] = len(formula_coords)

        if not diag["has_formulas"]:
            return diag

        missing_count = none_count

        if numeric_formula_count > 0 and not any_numeric_value:
            missing_count = max(missing_count, int(round(len(formula_coords) * 0.5)))

        diag["missing_cached_count"] = missing_count

        if len(formula_coords) <= 0:
            diag["risk_level"] = "ok"
            diag["confidence"] = 1.0
            return diag

        missing_ratio = missing_count / max(1, len(formula_coords))
        if missing_ratio > 0.7:
            diag["risk_level"] = "blocked"
            diag["reason"] = f"Most formulas ({missing_count}/{len(formula_coords)}) have no cached values. Open in Excel, calculate, save, and re-upload; or paste values."
            diag["confidence"] = 0.9
        elif missing_ratio > 0.3:
            diag["risk_level"] = "warning"
            diag["reason"] = f"Some formulas ({missing_count}/{len(formula_coords)}) may have missing cached values."
            diag["confidence"] = 0.7
        else:
            diag["risk_level"] = "ok"
            diag["reason"] = "Formulas detected with cached values present"
            diag["confidence"] = 0.95

        if empty_str_count > 0 and diag["risk_level"] == "ok":
            diag["confidence"] = min(diag["confidence"], 0.9)

        return diag

    def _detect_regions_2d(self, sheet_name: str, max_scan_rows: int = 500, max_scan_cols: int = 80) -> List[Dict[str, Any]]:
        ws = self.wb_values[sheet_name]

        grid: Dict[Tuple[int, int], bool] = {}
        max_row_seen = 0
        max_col_seen = 0

        for r_idx, row in enumerate(ws.iter_rows(values_only=True, max_row=max_scan_rows, max_col=max_scan_cols), start=1):
            for c_idx, cell in enumerate(row, start=1):
                if cell is None:
                    continue
                if isinstance(cell, str) and not cell.strip():
                    continue
                grid[(r_idx, c_idx)] = True
                max_row_seen = max(max_row_seen, r_idx)
                max_col_seen = max(max_col_seen, c_idx)

        if not grid:
            return []

        visited: Set[Tuple[int, int]] = set()
        regions: List[Dict[str, Any]] = []

        def flood_fill(sr: int, sc: int) -> Dict[str, Any]:
            stack = [(sr, sc)]
            cells: Set[Tuple[int, int]] = set()
            min_r = max_r = sr
            min_c = max_c = sc

            while stack:
                r, c = stack.pop()
                if (r, c) in visited or (r, c) not in grid:
                    continue
                visited.add((r, c))
                cells.add((r, c))
                min_r = min(min_r, r)
                max_r = max(max_r, r)
                min_c = min(min_c, c)
                max_c = max(max_c, c)

                for dr in (-1, 0, 1):
                    for dc in (-1, 0, 1):
                        if dr == 0 and dc == 0:
                            continue
                        for gap in (1, 2):
                            nr, nc = r + dr * gap, c + dc * gap
                            if (nr, nc) in grid and (nr, nc) not in visited:
                                stack.append((nr, nc))

            return {"bounds": (min_r, max_r, min_c, max_c), "cell_count": len(cells)}

        for (r, c) in list(grid.keys()):
            if (r, c) in visited:
                continue
            region = flood_fill(r, c)
            if region["cell_count"] < 9:
                continue
            regions.append(region)

        enriched: List[Dict[str, Any]] = []
        for region in regions:
            min_r, max_r, min_c, max_c = region["bounds"]

            # Header row: scan a small window of rows within the region and score.
            # Many real-world sheets have title rows above the actual table header.
            first_rows: List[List[Any]] = []
            scan_n = min(12, max(1, (max_r - min_r + 1)))
            for r in range(min_r, min(min_r + scan_n, max_r + 1)):
                row_data: List[Any] = []
                for c in range(min_c, max_c + 1):
                    row_data.append(ws.cell(r, c).value)
                first_rows.append(row_data)

            header_row_idx: Optional[int] = None
            if first_rows:
                scores: List[float] = []
                for row in first_rows:
                    text_count = sum(1 for v in row if isinstance(v, str) and v.strip())
                    num_count = sum(1 for v in row if isinstance(v, (int, float)))

                    # Prefer rows that look like headers for wide matrices, e.g. Jan-25..Dec-25.
                    period_hits = 0
                    keyword_hits = 0
                    for v in row:
                        if isinstance(v, str) and v.strip():
                            if self._is_period_header(v):
                                period_hits += 1
                            vv = v.strip().lower()
                            if vv in {"salesperson", "customer", "product", "region", "date", "total sales", "order", "order no"}:
                                keyword_hits += 1

                    scores.append(text_count + (period_hits * 3.0) + (keyword_hits * 2.0) - (num_count * 0.5))
                best_idx = scores.index(max(scores))
                header_row_idx = min_r + best_idx

            conf = min(region["cell_count"] / 50, 1.0)
            enriched.append(
                {
                    "bounds": region["bounds"],
                    "header_row": header_row_idx,
                    "data_start_row": (header_row_idx + 1) if header_row_idx else min_r,
                    "cell_count": region["cell_count"],
                    "confidence": conf,
                    "scan_limits": {"max_rows": max_row_seen, "max_cols": max_col_seen},
                }
            )

        enriched.sort(key=lambda r: r["confidence"], reverse=True)
        return enriched

    def _is_period_header(self, header: Any) -> bool:
        if header is None:
            return False
        if isinstance(header, datetime):
            return True
        if not isinstance(header, str):
            return False

        t = header.strip().lower()
        if not t:
            return False

        # Exclude summary columns like "12-month Total" / "YTD Total".
        if "total" in t or "subtotal" in t:
            return False

        # Strict patterns only (avoid treating arbitrary text containing 'jan' etc. as a period).
        # YYYY-MM or YYYY/M
        if re.match(r"^\d{4}[-/]\d{1,2}$", t):
            return True

        # Q1 2025 / Q1-25
        if re.match(r"^q[1-4]\s*[-/]?\s*\d{2,4}$", t):
            return True

        # FY2025 / FY 25
        if re.match(r"^fy\s*\d{2,4}$", t):
            return True

        # Jan-25 / January 2025
        m = re.match(r"^([a-z]{3,9})\s*[-/]?\s*(\d{2,4})$", t)
        if m:
            mon = m.group(1)[:3]
            if mon in {"jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"}:
                return True

        return False

    def _detect_column_roles(self, sheet_name: str, region: Dict[str, Any]) -> Dict[str, Any]:
        ws = self.wb_values[sheet_name]
        min_r, max_r, min_c, max_c = region["bounds"]
        data_start = int(region["data_start_row"])

        profiles: Dict[int, Dict[str, Any]] = {}

        for col_idx in range(min_c, max_c + 1):
            p: Dict[str, Any] = {
                "col_idx": col_idx,
                "text_count": 0,
                "numeric_count": 0,
                "date_count": 0,
                "distinct_text": set(),
                "total": 0,
            }

            for row_idx in range(data_start, min(data_start + 100, max_r + 1)):
                v = ws.cell(row_idx, col_idx).value
                if v is None:
                    continue
                if isinstance(v, str) and not v.strip():
                    continue
                p["total"] += 1
                if isinstance(v, str):
                    p["text_count"] += 1
                    if len(p["distinct_text"]) <= 500:
                        p["distinct_text"].add(v.strip())
                elif isinstance(v, (int, float)):
                    p["numeric_count"] += 1
                elif isinstance(v, datetime):
                    p["date_count"] += 1

            total = p["total"]
            p["text_ratio"] = (p["text_count"] / total) if total else 0.0
            p["numeric_ratio"] = (p["numeric_count"] / total) if total else 0.0
            p["date_ratio"] = (p["date_count"] / total) if total else 0.0
            p["cardinality"] = len(p["distinct_text"])
            profiles[col_idx] = p

        category_cols: List[int] = []
        value_cols: List[int] = []
        date_cols: List[int] = []

        for col_idx, p in profiles.items():
            if p["total"] < 3:
                continue
            if p["text_ratio"] >= 0.8 and 2 <= p["cardinality"] <= 50:
                category_cols.append(col_idx)
            elif p["numeric_ratio"] >= 0.7:
                value_cols.append(col_idx)
            elif p["date_ratio"] >= 0.5:
                date_cols.append(col_idx)

        period_headers: List[Dict[str, Any]] = []
        if region.get("header_row"):
            hr = int(region["header_row"])
            for col_idx in range(min_c, max_c + 1):
                hv = ws.cell(hr, col_idx).value
                if self._is_period_header(hv):
                    period_headers.append({"col_idx": col_idx, "label": hv})

        return {
            "category_columns": category_cols,
            "value_columns": value_cols,
            "date_columns": date_cols,
            "period_headers": period_headers,
            "column_profiles": {
                k: {kk: vv for kk, vv in v.items() if kk != "distinct_text"} | {"cardinality": v.get("cardinality", 0)}
                for k, v in profiles.items()
            },
        }

    def _aggregate(self, sheet_name: str, region: Dict[str, Any], roles: Dict[str, Any]) -> Dict[str, Any]:
        ws = self.wb_values[sheet_name]
        min_r, max_r, min_c, max_c = region["bounds"]
        data_start = int(region["data_start_row"])

        category_totals: Dict[str, float] = defaultdict(float)
        time_series: List[Dict[str, Any]] = []

        category_col: Optional[int] = None
        if roles.get("category_columns"):
            category_col = min(
                roles["category_columns"],
                key=lambda c: roles.get("column_profiles", {}).get(c, {}).get("cardinality", 10**9),
            )

        period_headers = roles.get("period_headers") or []
        period_cols = [p["col_idx"] for p in period_headers if isinstance(p, dict) and p.get("col_idx")]

        if period_cols:
            for row_idx in range(data_start, max_r + 1):
                cat = None
                if category_col:
                    cv = ws.cell(row_idx, category_col).value
                    if cv is not None and not (isinstance(cv, str) and not cv.strip()):
                        cat = str(cv).strip()

                # Avoid treating the summary Total row as a category.
                if cat and cat.strip().lower() == "total":
                    continue

                row_total = 0.0
                row_has_any = False
                for pc in period_cols:
                    v = ws.cell(row_idx, pc).value
                    if isinstance(v, (int, float)):
                        row_total += float(v)
                        row_has_any = True

                if cat and row_has_any:
                    category_totals[cat] += row_total

            for p in period_headers:
                pc = p.get("col_idx")
                label = p.get("label")
                if not pc:
                    continue
                total = 0.0
                ok = False
                for row_idx in range(data_start, max_r + 1):
                    v = ws.cell(row_idx, pc).value
                    if isinstance(v, (int, float)):
                        total += float(v)
                        ok = True
                if ok:
                    time_series.append({"period": label, "value": total})

            return {
                "category_totals": dict(category_totals),
                "time_series": time_series,
                "metadata": {
                    "region": f"R{min_r}C{min_c}:R{max_r}C{max_c}",
                    "row_count": max(0, max_r - data_start + 1),
                    "method": "wide_period_headers",
                },
            }

        date_cols = roles.get("date_columns") or []
        value_cols = roles.get("value_columns") or []
        date_col = date_cols[0] if date_cols else None
        value_col = value_cols[0] if value_cols else None

        if date_col and value_col:
            bucket: Dict[str, float] = defaultdict(float)
            for row_idx in range(data_start, max_r + 1):
                dv = ws.cell(row_idx, date_col).value
                vv = ws.cell(row_idx, value_col).value
                if dv is None or not isinstance(vv, (int, float)):
                    continue
                if isinstance(dv, datetime):
                    key = dv.date().isoformat()
                else:
                    key = str(dv).strip()
                    if not key:
                        continue
                bucket[key] += float(vv)

            for k in sorted(bucket.keys()):
                time_series.append({"period": k, "value": bucket[k]})

        if category_col and value_cols:
            for row_idx in range(data_start, max_r + 1):
                cv = ws.cell(row_idx, category_col).value
                if cv is None:
                    continue
                cat = str(cv).strip()
                if not cat:
                    continue
                if cat.strip().lower() == "total":
                    continue
                total = 0.0
                ok = False
                for vc in value_cols:
                    v = ws.cell(row_idx, vc).value
                    if isinstance(v, (int, float)):
                        total += float(v)
                        ok = True
                if ok:
                    category_totals[cat] += total

        return {
            "category_totals": dict(category_totals),
            "time_series": time_series,
            "metadata": {
                "region": f"R{min_r}C{min_c}:R{max_r}C{max_c}",
                "row_count": max(0, max_r - data_start + 1),
                "method": "smart_roles",
            },
        }

    def process_universal(self) -> Dict[str, Any]:
        results: Dict[str, Any] = {
            "filename": self.filename,
            "file_size_mb": round(self.file_size_mb, 2),
            "diagnostics": {"sheets": []},
            "charts": [],
            "status": "success",
        }

        for sheet_name in self.wb_values.sheetnames:
            diag = self.diagnose_sheet_formulas(sheet_name)
            results["diagnostics"]["sheets"].append(diag)

            if diag.get("risk_level") == "blocked":
                continue

            regions = self._detect_regions_2d(sheet_name)
            if not regions:
                continue

            best = regions[0]
            roles = self._detect_column_roles(sheet_name, best)
            aggs = self._aggregate(sheet_name, best, roles)

            cat_totals = aggs.get("category_totals") or {}
            if cat_totals:
                sorted_cats = sorted(cat_totals.items(), key=lambda x: abs(x[1]), reverse=True)[:10]
                results["charts"].append(
                    {
                        "type": "bar",
                        "title": f"Top Categories - {sheet_name}",
                        "data": {"labels": [c for c, _v in sorted_cats], "values": [v for _c, v in sorted_cats]},
                        "provenance": {
                            "sheet": sheet_name,
                            "region": aggs.get("metadata", {}).get("region"),
                            "method": aggs.get("metadata", {}).get("method"),
                            "confidence": best.get("confidence"),
                            "category_column": (roles.get("category_columns") or [None])[0],
                        },
                    }
                )

            ts = aggs.get("time_series") or []
            if isinstance(ts, list) and len(ts) >= 3:
                results["charts"].append(
                    {
                        "type": "line",
                        "title": f"Trend - {sheet_name}",
                        "data": {"labels": [p.get("period") for p in ts], "values": [p.get("value") for p in ts]},
                        "provenance": {
                            "sheet": sheet_name,
                            "region": aggs.get("metadata", {}).get("region"),
                            "method": aggs.get("metadata", {}).get("method"),
                            "confidence": best.get("confidence"),
                        },
                    }
                )

        results["total_charts"] = len(results["charts"])

        blocked_count = sum(1 for s in (results.get("diagnostics", {}).get("sheets") or []) if s.get("risk_level") == "blocked")
        if blocked_count >= len(results["diagnostics"]["sheets"]) and len(results["diagnostics"]["sheets"]) > 0:
            results["status"] = "blocked"
            results["message"] = "All sheets have missing formula cached values. Please recalculate in Excel."
        elif blocked_count > 0:
            results["status"] = "partial"
            results["message"] = f"{blocked_count} sheet(s) blocked due to missing cached values."

        return results
