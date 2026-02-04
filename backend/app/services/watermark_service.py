import io
from typing import Optional


def should_apply_watermark(plan: Optional[str]) -> bool:
    p = (plan or "").strip().lower()
    return p in ("", "free", "trial")


def watermark_pdf_bytes(pdf_bytes: bytes, text: str = "meldra.ai") -> bytes:
    """Add a visible diagonal watermark on each page of a PDF."""
    try:
        from PyPDF2 import PdfReader, PdfWriter
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.colors import Color
    except Exception:
        return pdf_bytes

    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        writer = PdfWriter()

        for page in reader.pages:
            w = float(page.mediabox.width)
            h = float(page.mediabox.height)

            packet = io.BytesIO()
            c = canvas.Canvas(packet, pagesize=(w, h))
            c.saveState()
            c.translate(w / 2.0, h / 2.0)
            c.rotate(35)
            c.setFillColor(Color(0.2, 0.2, 0.2, alpha=0.15))
            c.setFont("Helvetica-Bold", max(min(w, h) / 10.0, 36))
            c.drawCentredString(0, 0, text)
            c.restoreState()
            c.showPage()
            c.save()
            packet.seek(0)

            wm_pdf = PdfReader(packet)
            wm_page = wm_pdf.pages[0]
            page.merge_page(wm_page)
            writer.add_page(page)

        out = io.BytesIO()
        writer.write(out)
        out.seek(0)
        return out.read()
    except Exception:
        return pdf_bytes


def watermark_pptx_bytes(pptx_bytes: bytes, text: str = "meldra.ai") -> bytes:
    """Add a small footer watermark to every slide."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except Exception:
        return pptx_bytes

    try:
        prs = Presentation(io.BytesIO(pptx_bytes))
        for slide in prs.slides:
            # Bottom-right small watermark
            tb = slide.shapes.add_textbox(Inches(8.3), Inches(6.95), Inches(1.5), Inches(0.3))
            tf = tb.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = text
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(120, 120, 120)
            # transparency not directly supported; keep subtle via color

        out = io.BytesIO()
        prs.save(out)
        out.seek(0)
        return out.read()
    except Exception:
        return pptx_bytes


def watermark_xlsx_bytes(xlsx_bytes: bytes, text: str = "meldra.ai") -> bytes:
    """Add a simple watermark header cell to the first worksheet (reliable across viewers)."""
    try:
        import openpyxl
    except Exception:
        return xlsx_bytes

    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.worksheets[0] if wb.worksheets else wb.create_sheet("Sheet1")

        # Insert a top row and write watermark
        ws.insert_rows(1)
        ws["A1"].value = f"Generated with {text} (Free plan)"
        try:
            ws["A1"].font = openpyxl.styles.Font(bold=True, color="808080")
        except Exception:
            pass

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out.read()
    except Exception:
        return xlsx_bytes
