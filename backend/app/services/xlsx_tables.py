"""
Find the tables on a worksheet and show them as the sheet shows them.

A "table" is a block of filled cells; a block may contain single blank rows (e.g. a blank line
before a Totals row) but is separated from other blocks by two or more blank rows or columns.
A text-only first row spanning the block (e.g. "World-wide Sales - Millions of Dollars") is
treated as the table's title. Values are formatted with the cell's own Excel number format, so
currency symbols, thousands separators, percentages and dates look the way they do in Excel.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import List, Optional, Tuple

MAX_SCAN_ROWS = 5000
MAX_SCAN_COLS = 80


@dataclass
class TableCell:
    text: str
    numeric: bool = False
    bold: bool = False


@dataclass
class SheetTable:
    sheet: str
    title: str
    rows: List[List[TableCell]]  # first row is the header
    order: Tuple[int, int]  # (row, col) of the top-left cell, 0-based, for reading order
    total_rows: int = 0  # data rows in the sheet (rows may be cut for very long tables)
    extra: dict = field(default_factory=dict)


# ---------------------------------------------------------------- number formats

_DATE_TOKENS = [
    ("yyyy", "%Y"), ("yy", "%y"), ("mmmmm", "%b"), ("mmmm", "%B"), ("mmm", "%b"),
    ("dddd", "%A"), ("ddd", "%a"), ("dd", "%d"), ("hh", "%H"), ("ss", "%S"),
]


def _clean_section(fmt: str) -> str:
    fmt = re.sub(r"\[\$([^\]-]*)(-[^\]]*)?\]", lambda m: m.group(1), fmt)  # [$£-809] -> £
    fmt = re.sub(r"\[[^\]]*\]", "", fmt)  # colours, conditions
    fmt = re.sub(r"_.", "", fmt)  # padding
    fmt = fmt.replace("*", "").replace("\\", "")
    return fmt


def _format_date(value, fmt: str) -> str:
    f = _clean_section(fmt.split(";")[0]).lower()
    if not re.search(r"[dmyhs]", f):
        f = "dd mmm yyyy"
    out, i = "", 0
    has_time = "h" in f
    while i < len(f):
        for tok, rep in _DATE_TOKENS:
            if f.startswith(tok, i):
                out += rep
                i += len(tok)
                break
        else:
            ch = f[i]
            if ch == "m":
                # "mm"/"m" is minutes next to hours/seconds, otherwise the month
                run = 2 if f.startswith("mm", i) else 1
                minutes = has_time and (re.search(r"h[^dmy]*$", f[:i]) or re.match(r"^[^dmy]*s", f[i + run:]))
                out += "%M" if minutes else ("%m" if run == 2 else "%-m")
                i += run
                continue
            if ch == "d":
                out += "%-d"
            elif ch == "h":
                out += "%-H"
            elif ch == "s":
                out += "%-S"
            elif ch == '"':
                j = f.find('"', i + 1)
                out += f[i + 1:j] if j > i else ""
                i = j + 1 if j > i else len(f)
                continue
            elif f.startswith("am/pm", i):
                out += "%p"
                i += 5
                continue
            else:
                out += ch
            i += 1
    if isinstance(value, time):
        value = datetime.combine(date(1900, 1, 1), value)
    try:
        return value.strftime(out)
    except ValueError:
        return value.strftime(out.replace("%-", "%"))


def format_number(value: float, fmt: str) -> str:
    """Format a number with an Excel number format (the common cases: decimals, thousands, %, currency, negatives)."""
    fmt = (fmt or "General").strip()
    if fmt.lower() == "general" or fmt == "@":
        if float(value).is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.10g}"
    sections = fmt.split(";")
    section = sections[0]
    negative_style = None
    if value < 0 and len(sections) > 1 and sections[1].strip():
        section = sections[1]
        negative_style = "explicit"
        value = abs(value)
    section = _clean_section(section)
    literals = re.findall(r'"([^"]*)"', section)
    body = re.sub(r'"[^"]*"', "\x00", section)
    percent = "%" in body
    if percent:
        value = value * 100
    m = re.search(r"[0#?,.]+", body)
    if not m:
        return f"{value:g}"
    pattern = m.group(0)
    decimals = len(pattern.split(".", 1)[1].replace(",", "")) if "." in pattern else 0
    thousands = "," in pattern.split(".")[0]
    scale_commas = len(pattern) - len(pattern.rstrip(","))  # "#,##0," divides by 1,000
    value = value / (1000 ** scale_commas)
    number = f"{value:,.{decimals}f}" if thousands else f"{value:.{decimals}f}"
    prefix, suffix = body[: m.start()], body[m.end():]
    for lit in literals:
        prefix = prefix.replace("\x00", lit, 1) if "\x00" in prefix else prefix
        suffix = suffix.replace("\x00", lit, 1) if "\x00" in suffix else suffix
    text = f"{prefix}{number}{suffix}".replace("\x00", "")
    if negative_style == "explicit":
        text = text.replace("-", "", 0)
    return text.strip()


def format_cell(value, fmt: str) -> Tuple[str, bool]:
    """(display text, is_numeric)"""
    if value is None:
        return "", False
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), False
    if isinstance(value, (datetime, date, time)):
        return _format_date(value, fmt or "dd mmm yyyy"), True
    if isinstance(value, (int, float)):
        try:
            return format_number(float(value), fmt), True
        except Exception:
            return str(value), True
    return str(value).strip(), False


# ---------------------------------------------------------------- finding tables

def _blocks(filled: List[List[bool]]) -> List[Tuple[int, int, int, int]]:
    """(r0, r1, c0, c1) inclusive blocks of filled cells; one blank row inside a block is allowed."""
    n_rows = len(filled)
    row_has = [any(r) for r in filled]
    groups, start, blank = [], None, 0
    for i in range(n_rows + 2):
        has = i < n_rows and row_has[i]
        if has:
            if start is None:
                start = i
            blank = 0
            end = i
        elif start is not None:
            blank += 1
            if blank >= 2 or i >= n_rows:
                groups.append((start, end))
                start, blank = None, 0
    blocks = []
    for r0, r1 in groups:
        cols = [any(filled[r][c] for r in range(r0, r1 + 1)) for c in range(len(filled[0]))]
        c_start, c_blank = None, 0
        for c in range(len(cols) + 2):
            has = c < len(cols) and cols[c]
            if has:
                if c_start is None:
                    c_start = c
                c_blank = 0
                c_end = c
            elif c_start is not None:
                c_blank += 1
                if c_blank >= 2 or c >= len(cols):
                    blocks.append((r0, r1, c_start, c_end))
                    c_start, c_blank = None, 0
    return blocks


def find_tables(ws, sheet_name: str, max_rows: int = 200) -> List[SheetTable]:
    """Tables on an openpyxl worksheet (loaded with data_only=True, not read-only)."""
    n_rows = min(ws.max_row or 0, MAX_SCAN_ROWS)
    n_cols = min(ws.max_column or 0, MAX_SCAN_COLS)
    if n_rows == 0 or n_cols == 0:
        return []
    hidden_rows = {i for i, d in ws.row_dimensions.items() if d.hidden}
    hidden_cols = set()
    from openpyxl.utils import column_index_from_string

    for key, d in ws.column_dimensions.items():
        if d.hidden:
            lo = column_index_from_string(key)
            hi = d.max or lo
            hidden_cols.update(range(max(lo, d.min or lo), hi + 1))
    merged_into = {}
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if (r, c) != (mr.min_row, mr.min_col):
                    merged_into[(r, c)] = (mr.min_row, mr.min_col)

    cells = [[None] * n_cols for _ in range(n_rows)]
    filled = [[False] * n_cols for _ in range(n_rows)]
    for row in ws.iter_rows(min_row=1, max_row=n_rows, max_col=n_cols):
        for cell in row:
            r, c = cell.row - 1, cell.column - 1
            if cell.row in hidden_rows or cell.column in hidden_cols:
                continue
            if cell.value is not None and str(cell.value).strip() != "":
                cells[r][c] = cell
                filled[r][c] = True
    for (r, c), (tr, tc) in merged_into.items():  # a merged range counts as filled across its span
        if r <= n_rows and c <= n_cols and filled[tr - 1][tc - 1]:
            filled[r - 1][c - 1] = True

    tables = []
    for r0, r1, c0, c1 in _blocks(filled):
        # Title lines: rows at the top holding a single piece of text (e.g. a heading and a subtitle).
        title_lines = []
        while r0 < r1:
            texts = [cells[r0][c] for c in range(c0, c1 + 1) if cells[r0][c] is not None]
            if len(texts) != 1 or not isinstance(texts[0].value, str):
                break
            title_lines.append(str(texts[0].value).strip())
            r0 += 1
            while r0 <= r1 and not any(filled[r0][c] for c in range(c0, c1 + 1)):
                r0 += 1
        title = " — ".join(title_lines)
        # drop fully blank rows inside the block (the single allowed gaps)
        row_idx = [r for r in range(r0, r1 + 1) if any(cells[r][c] is not None for c in range(c0, c1 + 1))]
        col_idx = [c for c in range(c0, c1 + 1) if any(cells[r][c] is not None for r in row_idx)]
        # Columns that are nearly empty in a long table are captions or helper cells, not data:
        # a caption in the header row becomes the title if there is none.
        data_rows = row_idx[1:]
        if len(data_rows) >= 5:
            keep = []
            for c in col_idx:
                used = sum(1 for r in data_rows if cells[r][c] is not None)
                if used >= max(2, len(data_rows) // 10):
                    keep.append(c)
                    continue
                head = cells[row_idx[0]][c]
                if not title and head is not None and isinstance(head.value, str) and len(head.value.strip()) > 20:
                    title = head.value.strip()
            col_idx = keep
        if len(row_idx) < 2 or len(col_idx) < 2:
            continue  # a note or a single value is not a table
        rows = []
        for r in row_idx[: max_rows + 1]:
            out = []
            for c in col_idx:
                cell = cells[r][c]
                if cell is None:
                    out.append(TableCell(""))
                    continue
                text, numeric = format_cell(cell.value, cell.number_format)
                out.append(TableCell(text, numeric, bool(cell.font and cell.font.b)))
            rows.append(out)
        tables.append(SheetTable(sheet_name, title, rows, (row_idx[0], col_idx[0]), total_rows=len(row_idx) - 1))
    return tables
