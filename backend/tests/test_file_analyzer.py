"""File check (Workbench step 3): multi-tab workbooks, titles above tables, empty and chart-only tabs."""
import asyncio
import io

import openpyxl
import pytest

from app.services import file_analyzer
from app.services.file_analyzer import FileAnalyzerService, _table_from_rows


@pytest.fixture(autouse=True)
def _no_ai(monkeypatch):
    async def fake_llm(*args, **kwargs):
        raise RuntimeError("AI not available in tests")
    monkeypatch.setattr(file_analyzer, "invoke_llm", fake_llm)


def _workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Monthly sales by region"])
    ws.append([])
    ws.append(["Month", "North", "South"])
    for m, n, s in [("Jan", 100, 200), ("Feb", 150, 210), ("Mar", 170, 190)]:
        ws.append([m, n, s])
    pl = wb.create_sheet("P&L")
    pl.append(["Line item", "Q1", "Q2"])
    pl.append(["Revenue", 1200, 1350])
    pl.append(["Net profit", 400, 470])
    wb.create_sheet("Chart only")
    wb.create_sheet("Headers only").append(["A", "B", "C"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_multi_tab_workbook_with_empty_tabs_is_analysed():
    # This workbook used to fail with "division by zero".
    r = asyncio.run(FileAnalyzerService().analyze_excel_file(_workbook(), "book.xlsx"))
    sheets = {s["name"]: s for s in r["sheets"]}
    assert [c["name"] for c in sheets["Sales"]["columns"]] == ["Month", "North", "South"]
    assert sheets["Sales"]["row_count"] == 3
    assert sheets["P&L"]["row_count"] == 2
    for name in ("Chart only", "Headers only"):
        assert sheets[name]["empty"] is True and sheets[name]["row_count"] == 0 and sheets[name]["note"]
    assert r["overall_summary"]["total_rows"] == 5


def test_one_bad_tab_does_not_fail_the_file(monkeypatch):
    real = FileAnalyzerService._analyze_sheet

    async def flaky(self, sheet_data, filename):
        if sheet_data["name"] == "P&L":
            raise ValueError("boom")
        return await real(self, sheet_data, filename)

    monkeypatch.setattr(FileAnalyzerService, "_analyze_sheet", flaky)
    r = asyncio.run(FileAnalyzerService().analyze_excel_file(_workbook(), "book.xlsx"))
    sheets = {s["name"]: s for s in r["sheets"]}
    assert "boom" in sheets["P&L"]["note"] and sheets["Sales"]["row_count"] == 3


def test_header_row_is_found_below_titles_and_blank_rows():
    headers, rows = _table_from_rows(
        [["Report title", None, None], [None, None, None], ["Name", "Dept", "Pay"], ["Asha", "HR", 10], [None, None, None], ["Ben", "IT", 20]],
        100,
    )
    assert headers == ["Name", "Dept", "Pay"]
    assert rows == [["Asha", "HR", 10], ["Ben", "IT", 20]]
    assert _table_from_rows([], 100) == ([], [])
