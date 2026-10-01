"""End-to-end checks of the file tools through the real API: upload a file, check what comes back."""
import io
import json
import os
import secrets
import tempfile
import zipfile

import pytest

os.environ.setdefault("DATABASE_URL", f"sqlite:///{os.path.join(tempfile.gettempdir(), 'meldra_connector_endpoints.db')}")
main = pytest.importorskip("app.main", reason="full backend dependencies not installed")
from fastapi.testclient import TestClient  # noqa: E402

from app.database import SessionLocal, User  # noqa: E402
from app.services.zip_processor import ZipProcessorService  # noqa: E402

main.init_db()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def client():
    email = f"files-{secrets.token_hex(4)}@example.com"
    db = SessionLocal()
    try:
        db.add(User(email=email, full_name="File Tools", hashed_password="x", is_verified=True))
        db.commit()
    finally:
        db.close()
    main.app.dependency_overrides[main.get_current_user] = lambda: {"email": email}
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def _xlsx(rows, title="Sheet1", chart=False):
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    if chart:
        c = BarChart()
        c.title = "Sales by region"
        c.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(rows)), titles_from_data=True)
        c.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(rows)))
        ws.add_chart(c, "E2")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf(pages):
    from PyPDF2 import PdfWriter

    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(200, 200)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def _zip(entries):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, data in entries.items():
            z.writestr(zipfile.ZipInfo(name, date_time=(2024, 3, 1, 12, 0, 0)), data)
    return out.getvalue()


# ---- Filename cleaning -----------------------------------------------------------------------

def test_sanitize_filename_rules():
    s = ZipProcessorService()
    # Ranges in the allowed characters, extension kept, uppercase lowered instead of replaced.
    assert s.sanitize_filename("Q1 Report (Final).PDF", allowed_chars="a-z0-9-_", remove_spaces=True) == "q1report_final.pdf"
    assert s.sanitize_filename("Übersicht März.xlsx", allowed_chars="a-zA-Z0-9-", replace_char="-") == "Ubersicht-Marz.xlsx"
    assert s.sanitize_filename("Übersicht.xlsx", language_replacements={"Ü": "Ue"}) == "Uebersicht.xlsx"
    assert s.sanitize_filename("报告.pdf", allowed_chars="a-z") == "file.pdf"
    assert s.sanitize_filename("a" * 50 + ".xlsx", max_length=20) == "a" * 15 + ".xlsx"
    assert s.expand_char_set("a-c-") == {"a", "b", "c", "-"}


def test_process_zip_uses_form_options_keeps_folders_and_never_overwrites(client):
    data = _zip({"Q1 Übersicht/März Bericht.txt": "hello", "Résumé.pdf": "1", "Resume.pdf": "2", "notes..v2.txt": "n"})
    options = {"allowed_chars": "a-zA-Z0-9-_", "replace_char": "-", "languages": ["german"]}
    r = client.post("/api/files/process-zip", files={"file": ("Batch.ZIP", data, "application/zip")}, data={"options": json.dumps(options)})
    assert r.status_code == 200, r.text
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = sorted(z.namelist())
    assert names == ["Q1-Uebersicht/Maerz-Bericht.txt", "Resume-2.pdf", "Resume.pdf", "notes.v2.txt"]
    assert z.read("Q1-Uebersicht/Maerz-Bericht.txt") == b"hello"
    assert z.getinfo("Resume.pdf").date_time == (2024, 3, 1, 12, 0, 0)


def test_process_zip_rejects_bad_input_with_a_clear_message(client):
    zip_ct = "application/zip"
    assert client.post("/api/files/process-zip", files={"file": ("x.zip", _zip({"a.txt": "1"}), zip_ct)}, data={"options": "{bad"}).status_code == 400
    unsafe = client.post("/api/files/process-zip", files={"file": ("x.zip", _zip({"../evil.txt": "1"}), zip_ct)})
    assert unsafe.status_code == 400 and "unsafe" in unsafe.json()["detail"]
    assert client.post("/api/files/process-zip", files={"file": ("x.zip", b"not a zip", zip_ct)}).status_code == 400
    assert client.post("/api/files/process-zip", files={"file": ("x.txt", b"hi", "text/plain")}).status_code == 400


# ---- Excel to PowerPoint ---------------------------------------------------------------------

def test_excel_to_ppt_with_branding_keeps_native_charts(client):
    from pptx import Presentation

    data = _xlsx([["Region", "Sales"], ["North", 120], ["South", 90], ["East", 150]], chart=True)
    r = client.post(
        "/api/files/excel-to-ppt",
        files={"file": ("sales.xlsx", data, XLSX)},
        data={"theme": "dark", "brand_color": "#0A7E8C", "company": "Acme Ltd"},
    )
    assert r.status_code == 200, r.text
    deck = Presentation(io.BytesIO(r.content))
    shapes = [sh for slide in deck.slides for sh in slide.shapes]
    assert any(sh.has_chart for sh in shapes), "chart should stay an editable PowerPoint chart"
    assert any(sh.has_table for sh in shapes)
    text = " ".join(sh.text_frame.text for sh in shapes if sh.has_text_frame)
    assert "Acme Ltd" in text


def test_csv_to_ppt(client):
    from pptx import Presentation

    r = client.post("/api/files/excel-to-ppt", files={"file": ("data.csv", b"name,value\na,1\nb,2\n", "text/csv")})
    assert r.status_code == 200, r.text
    assert any(sh.has_table for slide in Presentation(io.BytesIO(r.content)).slides for sh in slide.shapes)


# ---- Spreadsheet tools -----------------------------------------------------------------------

def test_analyze_reads_formatted_numbers(client):
    data = _xlsx([["Invoice", "Amount"], ["A1", "1,200.50"], ["A2", "$300"], ["A3", "(45)"]])
    r = client.post("/api/files/analyze", files={"file": ("d.xlsx", data, XLSX)})
    assert r.status_code == 200, r.text
    amount = next(c for c in r.json()["sheets"][0]["columns"] if c["name"] == "Amount")
    assert amount["type"] == "numeric" and amount["max"] == 1200.5 and amount["min"] == -45


def test_standardize_and_reconcile(client):
    left = _xlsx([["Invoice", "Amount"], ["A1", "1,200.50"], ["A2", "300"], ["A2", "300"]])
    right = _xlsx([["Ref", "Total"], ["A1", 1200.5], ["A3", 50]])

    std = client.post("/api/files/standardize-preview", files={"file": ("l.xlsx", left, XLSX)})
    assert std.status_code == 200 and std.json()["duplicate_rows_removed"] == 1
    assert client.post("/api/files/standardize", files={"file": ("l.xlsx", left, XLSX)}).content[:2] == b"PK"

    form = {"left_key_col": "Invoice", "right_key_col": "Ref", "left_amount_col": "Amount", "right_amount_col": "Total"}
    files = {"left_file": ("l.xlsx", left, XLSX), "right_file": ("r.xlsx", right, XLSX)}
    rec = client.post("/api/files/reconcile-preview", files=files, data=form)
    assert rec.status_code == 200, rec.text
    assert rec.json()["counts"]["matched"] == 1 and rec.json()["counts"]["missing_on_left"] == 1
    assert client.post("/api/files/reconcile", files=files, data=form).content[:2] == b"PK"


# ---- PDF and document conversions ------------------------------------------------------------

def test_pdf_merge_and_split(client):
    from PyPDF2 import PdfReader

    merged = client.post(
        "/api/pdf/merge",
        files=[("files", ("a.pdf", _pdf(1), "application/pdf")), ("files", ("b.pdf", _pdf(2), "application/pdf"))],
    )
    assert merged.status_code == 200 and len(PdfReader(io.BytesIO(merged.content)).pages) == 3
    split = client.post("/api/pdf/split", files={"file": ("a.pdf", _pdf(3), "application/pdf")}, data={"page_ranges": "2-3"})
    assert split.status_code == 200, split.text


def test_document_conversions(client):
    import docx

    d = docx.Document()
    d.add_paragraph("Quarterly numbers")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Month", "Sales", "Jan", "10"
    buf = io.BytesIO()
    d.save(buf)

    for endpoint, name, body in (("doc-to-xls", "a.docx", buf.getvalue()), ("pdf-to-xls", "a.pdf", _pdf(1)), ("pdf-to-doc", "a.pdf", _pdf(1))):
        r = client.post(f"/api/convert/{endpoint}", files={"file": (name, body, "application/octet-stream")})
        assert r.status_code == 200, f"{endpoint}: {r.text}"
        assert r.content[:2] == b"PK", endpoint


def test_memory_is_released_after_a_file_request(client, monkeypatch):
    calls = []
    monkeypatch.setattr(main, "release_memory", lambda: calls.append(1))
    r = client.post("/api/files/excel-to-ppt", files={"file": ("data.csv", b"name,value\na,1\nb,2\n", "text/csv")})
    assert r.status_code == 200
    assert calls == [1]
    client.get("/api/auth/devices")  # ordinary requests don't trigger it
    assert calls == [1]


def test_upload_limit_switch(client, monkeypatch):
    # Plan limits by default (a new account is on the free plan).
    monkeypatch.delenv("UPLOAD_LIMIT_MB", raising=False)
    assert client.get("/api/subscriptions/me").json()["max_upload_mb"] == 10

    # A number applies to everyone: a 2 MB ZIP is refused at 1 MB.
    monkeypatch.setenv("UPLOAD_LIMIT_MB", "1")
    assert client.get("/api/subscriptions/me").json()["max_upload_mb"] == 1
    big = _zip({f"f{i}.bin": secrets.token_bytes(1024) for i in range(2100)})
    r = client.post("/api/files/process-zip", files={"file": ("big.zip", big, "application/zip")})
    assert r.status_code == 413

    # "off" removes the limit (for load testing) and the website is told there is none.
    monkeypatch.setenv("UPLOAD_LIMIT_MB", "off")
    assert client.get("/api/subscriptions/me").json()["max_upload_mb"] is None
    r = client.post("/api/files/process-zip", files={"file": ("big.zip", big, "application/zip")})
    assert r.status_code == 200, r.text
