"""
Meldra lakehouse: ingest (CSV / Excel / Parquet / JSON rows) into Apache Iceberg,
typing and profiling, tenant isolation, and server-side aggregation with the
browser engine's semantics.

Runs against an Iceberg SQL catalog in a temp directory by default. Set
POLARIS_TEST_URI (plus POLARIS_TEST_CREDENTIAL and POLARIS_TEST_WAREHOUSE) to run
the same tests against a real Apache Polaris server.
"""
import csv
import os
import uuid

import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from app.services.lakehouse import config, query, store

POLARIS = os.environ.get("POLARIS_TEST_URI") if os.environ.get("POLARIS_TEST_CREDENTIAL") else None  # both needed


@pytest.fixture(params=["sql"] + (["polaris"] if POLARIS else []))
def lake(request, tmp_path, monkeypatch):
    if request.param == "sql":
        monkeypatch.setenv("LAKEHOUSE_CATALOG", "sql")
        monkeypatch.setenv("LAKEHOUSE_WAREHOUSE", f"file://{tmp_path}/wh")
        monkeypatch.delenv("LAKEHOUSE_CATALOG_URI", raising=False)
    else:
        monkeypatch.setenv("LAKEHOUSE_CATALOG", "rest")
        monkeypatch.setenv("POLARIS_URI", POLARIS)
        monkeypatch.setenv("POLARIS_CREDENTIAL", os.environ["POLARIS_TEST_CREDENTIAL"])
        monkeypatch.setenv("POLARIS_WAREHOUSE", os.environ.get("POLARIS_TEST_WAREHOUSE", "meldra"))
        monkeypatch.setenv("POLARIS_ACCESS_DELEGATION", os.environ.get("POLARIS_TEST_DELEGATION", "none"))
    monkeypatch.setenv("LAKEHOUSE_NAMESPACE_PREFIX", f"test{uuid.uuid4().hex[:6]}")
    store.reset_catalog()
    yield request.param
    for tenant in ("a@example.com", "b@example.com"):
        try:
            store.drop_all(tenant)
        except Exception:
            pass
    store.reset_catalog()


def write_csv(path, header, rows, delimiter=",", bom=False):
    with open(path, "w", newline="", encoding="utf-8-sig" if bom else "utf-8") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow(header)
        w.writerows(rows)
    return str(path)


def by_group(result, split=False):
    return {((r["g"], r["s"]) if split else r["g"]): r for r in result["groups"]}


def test_disabled_without_configuration(monkeypatch):
    monkeypatch.delenv("LAKEHOUSE_CATALOG", raising=False)
    store.reset_catalog()
    assert config.describe() == {"enabled": False}
    with pytest.raises(store.LakehouseError, match="not configured"):
        store.get_catalog()


def test_csv_typing_matches_the_browser_rules(lake, tmp_path):
    path = write_csv(tmp_path / "hr.csv", ["Emp ID", "Dept", "Hire Date", "Salary (£)", "Year", "Adj"], [
        ["E1", "Sales", "01/03/2024", "£52,000", "2024", "(1,200.50)"],
        ["E2", "HR", "15/11/2023", "£48,500", "2023", "300"],
        ["E3", "sales", "12/25/2025", "£1,000.25", "2025", "-50"],
    ], bom=True)
    [src] = store.ingest_file("a@example.com", path, "hr.csv")
    cols = {c["key"]: c for c in src["columns"]}
    assert list(cols) == ["emp_id", "dept", "hire_date", "salary", "year", "adj"]
    assert cols["emp_id"]["role"] == "dimension" and cols["year"]["role"] == "dimension"
    assert cols["hire_date"]["type"] == "date"
    assert cols["salary"] == {**cols["salary"], "type": "number", "role": "measure", "unit": "money", "currency": "£"}
    assert src["row_count"] == 3 and src["month_range"] == ["2023-11", "2025-12"]
    assert src["values"]["dept"]["distinct"] == 3

    res = query.aggregate("a@example.com", [{"table": src["table"], "measure": "adj", "group_by": "month"}])[0]
    g = by_group(res)
    assert g["2024-03"]["sum"] == pytest.approx(-1200.5)  # (1,200.50) is negative
    assert g["2025-12"]["sum"] == -50  # 12/25/2025 read month-first because 25 can't be a month
    assert res["rows"] == 3


def test_semicolon_csv_duplicate_headers_and_blank_groups(lake, tmp_path):
    path = write_csv(tmp_path / "x.csv", ["Region", "Region", "Amount"], [["North", "N1", "10"], ["", "N2", "5"], ["North", "N3", "1,000"]], delimiter=";")
    [src] = store.ingest_file("a@example.com", path, "x.csv")
    assert [c["name"] for c in src["columns"]] == ["Region", "Region (2)", "Amount"]
    res = query.aggregate("a@example.com", [{"table": src["table"], "measure": "amount", "group_by": "region"}])[0]
    assert by_group(res)["North"]["sum"] == 1010
    assert by_group(res)["(blank)"]["sum"] == 5


def test_filters_follow_browser_semantics(lake, tmp_path):
    path = write_csv(tmp_path / "inv.csv", ["Customer", "Status", "Level", "Amount", "Date"], [
        ["Acme", "Overdue", "10", "100", "2026-01-05"],
        ["Acme", "paid", "9", "50", "2026-02-10"],
        ["Beta", "OVERDUE", "2", "30", "2026-03-01"],
        ["Beta", "Paid", "B", "20", "2025-12-31"],
    ])
    [src] = store.ingest_file("a@example.com", path, "inv.csv")
    t = src["table"]
    agg = lambda filters, gb="customer": by_group(query.aggregate("a@example.com", [{"table": t, "measure": "amount", "group_by": gb, "filters": filters}])[0])  # noqa: E731
    assert {k: v["sum"] for k, v in agg([{"dim": "status", "op": "eq", "value": "overdue"}]).items()} == {"Acme": 100, "Beta": 30}  # case-insensitive
    assert {k: v["sum"] for k, v in agg([{"dim": "status", "op": "neq", "value": "Overdue"}]).items()} == {"Acme": 50, "Beta": 20}
    # Numbers compare as numbers (10 >= 9 >= 2), text falls back to text ("B" >= "9").
    assert {k: v["sum"] for k, v in agg([{"dim": "level", "op": "gte", "value": "9"}]).items()} == {"Acme": 150, "Beta": 20}
    months = agg([{"dim": "month", "op": "gte", "value": "2026-01"}, {"dim": "month", "op": "lte", "value": "2026-02"}], "month")
    assert {k: v["sum"] for k, v in months.items()} == {"2026-01": 100, "2026-02": 50}


def test_split_and_lookup_across_tables(lake, tmp_path):
    emp = write_csv(tmp_path / "emp.csv", ["Emp ID", "Department", "Country"], [["E1", "Sales", "UK"], ["E2", "HR", "UK"], ["E3", "Sales", "DE"]])
    exp = write_csv(tmp_path / "exp.csv", ["Employee ID", "Category", "Amount"], [
        ["e1", "Travel", "10"], ["E1", "Meals", "5"], ["E2", "Travel", "7"], ["E3", "Travel", "3"], ["E9", "Travel", "100"]])
    [e] = store.ingest_file("a@example.com", emp, "emp.csv")
    [x] = store.ingest_file("a@example.com", exp, "exp.csv")
    links = query.suggest_links("a@example.com", [e["table"], x["table"]])
    assert links[0]["from"] == {"table": x["table"], "col": "employee_id"} and links[0]["to"] == {"table": e["table"], "col": "emp_id"}
    lk = [{"from_col": "employee_id", "table": e["table"], "to_col": "emp_id", "dims": ["department", "country"]}]
    res = query.aggregate("a@example.com", [{"table": x["table"], "measure": "amount", "group_by": "department", "split_by": "category", "lookups": lk}])[0]
    g = by_group(res, split=True)
    assert g[("Sales", "Travel")]["sum"] == 13 and g[("Sales", "Meals")]["sum"] == 5 and g[("HR", "Travel")]["sum"] == 7
    assert g[("(blank)", "Travel")]["sum"] == 100  # E9 has no employee record
    uk = query.aggregate("a@example.com", [{"table": x["table"], "measure": "amount", "group_by": "department", "lookups": lk,
                                            "filters": [{"dim": "country", "op": "eq", "value": "uk"}]}])[0]
    assert {k: v["sum"] for k, v in by_group(uk).items()} == {"Sales": 15, "HR": 7}


def test_excel_parquet_and_json_rows(lake, tmp_path):
    from openpyxl import Workbook
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Order", "Region", "Amount", "Booked"])
    ws.append(["O1", "North", 120.5, datetime(2026, 1, 3)])
    ws.append(["O2", "South", 80, datetime(2026, 2, 7)])
    wb.create_sheet("Empty")
    ws2 = wb.create_sheet("Targets")
    ws2.append(["Region", "Target"])
    ws2.append(["North", 100])
    wb.save(tmp_path / "book.xlsx")
    srcs = store.ingest_file("a@example.com", str(tmp_path / "book.xlsx"), "book.xlsx")
    assert [s["name"] for s in srcs] == ["book · Orders", "book · Targets"]
    orders = srcs[0]
    assert {c["key"]: c["type"] for c in orders["columns"]}["booked"] == "date"
    res = query.aggregate("a@example.com", [{"table": orders["table"], "measure": "amount", "group_by": "month"}])[0]
    assert {k: v["sum"] for k, v in by_group(res).items()} == {"2026-01": 120.5, "2026-02": 80}

    pq.write_table(pa.table({"sku": ["A", "B", "A"], "qty": [1, 2, 3], "price": [9.5, 1.0, 9.5]}), tmp_path / "p.parquet")
    [p] = store.ingest_file("a@example.com", str(tmp_path / "p.parquet"), "p.parquet")
    res = query.aggregate("a@example.com", [{"table": p["table"], "measure": "qty", "group_by": "sku"}])[0]
    assert {k: v["sum"] for k, v in by_group(res).items()} == {"A": 4, "B": 2}

    j = store.ingest_rows("a@example.com", "Workers", "SuccessFactors", ["userId", "department", "fte", "active"],
                          [{"userId": "E1", "department": "Sales", "fte": 1, "active": True}, {"userId": "E2", "department": "HR", "fte": 0.5, "active": False}],
                          "api", origin={"type": "api", "url": "https://api4.successfactors.com/x"})
    assert j["origin"]["url"].startswith("https://") and j["kind"] == "api"
    assert j["values"]["active"]["top"] == ["false", "true"]  # booleans written like the browser does
    res = query.aggregate("a@example.com", [{"table": j["table"], "measure": "fte", "group_by": "department"}])[0]
    assert {k: v["sum"] for k, v in by_group(res).items()} == {"Sales": 1, "HR": 0.5}


def test_tenants_are_isolated(lake, tmp_path):
    [src] = store.ingest_file("a@example.com", write_csv(tmp_path / "a.csv", ["K", "V"], [["x", "1"]]), "a.csv")
    assert store.namespace_for("a@example.com") != store.namespace_for("b@example.com")
    assert store.namespace_for("A@Example.com ") == store.namespace_for("a@example.com")
    assert store.list_tables("b@example.com") == []
    for call in (lambda: store.describe_table("b@example.com", src["table"]),
                 lambda: query.aggregate("b@example.com", [{"table": src["table"], "group_by": "k"}]),
                 lambda: store.drop_table("b@example.com", src["table"])):
        with pytest.raises(store.LakehouseError, match="Unknown table"):
            call()
    with pytest.raises(store.LakehouseError, match="Unknown table"):
        store.describe_table("a@example.com", "../etc/passwd")
    assert [t["table"] for t in store.list_tables("a@example.com")] == [src["table"]]


def test_unknown_columns_are_rejected_not_interpolated(lake, tmp_path):
    [src] = store.ingest_file("a@example.com", write_csv(tmp_path / "a.csv", ["K", "V"], [["x", "1"]]), "a.csv")
    for bad in ('k" FROM src; DROP TABLE x; --', "nope"):
        with pytest.raises(store.LakehouseError, match="Unknown"):
            query.aggregate("a@example.com", [{"table": src["table"], "group_by": bad}])
    # Filter values are parameters: quotes are just text.
    res = query.aggregate("a@example.com", [{"table": src["table"], "group_by": "k", "filters": [{"dim": "k", "op": "eq", "value": "x' OR '1'='1"}]}])[0]
    assert res["groups"] == []


def test_user_choices_are_saved_and_refresh_keeps_them(lake, tmp_path):
    [src] = store.ingest_file("a@example.com", write_csv(tmp_path / "c.csv", ["Dept", "Cost", "Code"], [["Sales", "10", "7"]]), "c.csv")
    upd = store.update_table("a@example.com", src["table"], {"system": "SAP S/4", "columns": [
        {"name": "Dept", "key": "department"}, {"name": "Code", "role": "ignore"}]})
    cols = {c["name"]: c for c in upd["columns"]}
    assert upd["system"] == "SAP S/4" and cols["Dept"]["key"] == "department" and cols["Code"]["role"] == "ignore"
    assert "department" in upd["values"]
    res = query.aggregate("a@example.com", [{"table": src["table"], "measure": "cost", "group_by": "department"}])[0]
    assert by_group(res)["Sales"]["sum"] == 10
    with pytest.raises(store.LakehouseError, match="no numbers"):
        store.update_table("a@example.com", src["table"], {"columns": [{"name": "Dept", "role": "measure"}]})

    fresh = store.ingest_rows("a@example.com", "c", "c", ["Dept", "Cost", "Code", "New"], [{"Dept": "HR", "Cost": "4", "Code": "1", "New": "y"}],
                              "database", replace_table=src["table"])
    cols = {c["name"]: c for c in fresh["columns"]}
    assert fresh["system"] == "SAP S/4" and cols["Dept"]["key"] == "department" and cols["Code"]["role"] == "ignore" and "New" in cols
    assert [t["table"] for t in store.list_tables("a@example.com")] == [fresh["table"]]


def test_delete_removes_every_file(lake, tmp_path):
    """Erasure: after a delete, none of the table's files exist (checked through the catalog's own FileIO)."""
    [src] = store.ingest_file("a@example.com", write_csv(tmp_path / "d.csv", ["K", "V"], [["x", "1"]]), "d.csv")
    store.update_table("a@example.com", src["table"], {"system": "Edited"})  # a second profile version and metadata file
    tbl = store.get_catalog().load_table((store.namespace_for("a@example.com"), src["table"]))
    files = store.table_files(tbl)
    assert any(f.endswith(".parquet") for f in files) and any("profile-" in f for f in files) and len(files) >= 5
    assert all(tbl.io.new_input(f).exists() for f in files)
    store.drop_table("a@example.com", src["table"])
    assert store.list_tables("a@example.com") == []
    assert [f for f in files if tbl.io.new_input(f).exists()] == []
    store.ingest_file("a@example.com", write_csv(tmp_path / "e.csv", ["K"], [["x"]]), "e.csv")
    assert store.drop_all("a@example.com") == 1 and store.list_tables("a@example.com") == []


def test_preview_and_bad_files(lake, tmp_path):
    [src] = store.ingest_file("a@example.com", write_csv(tmp_path / "f.csv", ["When", "Amount"], [["2026-01-02", "5"]]), "f.csv")
    pv = store.preview("a@example.com", src["table"])
    assert pv["columns"] == ["When", "Amount"] and pv["rows"] == [["2026-01-02", 5.0]]
    (tmp_path / "empty.csv").write_text("")
    with pytest.raises(store.LakehouseError):
        store.ingest_file("a@example.com", str(tmp_path / "empty.csv"), "empty.csv")
    with pytest.raises(store.LakehouseError, match="upload a"):
        store.ingest_file("a@example.com", str(tmp_path / "f.csv"), "f.exe")


def test_currency_from_column_names():
    from app.services.lakehouse.profile import currency_from_name
    assert [currency_from_name(n) for n in ["Budget (£)", "licence_cost_gbp", "Cost USD", "amount_eur", "Amount", "gbpx"]] == ["£", "£", "$", "€", None, None]
