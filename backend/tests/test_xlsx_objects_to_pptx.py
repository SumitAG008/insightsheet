"""Excel -> PowerPoint with charts copied as native charts and pictures at full quality."""
import io
import os
import tempfile
import zipfile

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from pptx import Presentation

from app.services.xlsx_objects_to_pptx import convert_workbook_objects


def _png(color, size=(240, 120)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    b.seek(0)
    return b


def _bar(ws, title, grouping="clustered"):
    c = BarChart()
    c.title = title
    c.grouping = grouping
    if grouping == "stacked":
        c.overlap = 100
    c.add_data(Reference(ws, min_col=2, min_row=1, max_col=3, max_row=7), titles_from_data=True)
    c.set_categories(Reference(ws, min_col=1, min_row=2, max_row=7))
    return c


def _workbook() -> bytes:
    wb = openpyxl.Workbook()
    rows = [["Month", "North", "South"]] + [[f"M{i}", i * 10, i * 7] for i in range(1, 7)]

    sales = wb.active
    sales.title = "Sales"
    for r in rows:
        sales.append(r)
    sales.add_chart(_bar(sales, "Sales by region"), "E2")
    sales.add_chart(_bar(sales, "Sales by region", "stacked"), "E20")  # same data, looks different
    sales.add_image(XLImage(_png("navy")), "N2")  # logo

    copy = wb.create_sheet("Copy")
    for r in rows:
        copy.append(r)
    copy.add_chart(_bar(copy, "Sales by region"), "E2")  # same chart, same numbers, other tab
    copy.add_image(XLImage(_png("navy")), "N2")  # the same logo again
    copy.add_image(XLImage(_png("orange", (400, 400))), "N20")  # a different picture

    empty = wb.create_sheet("Empty")
    empty.add_chart(BarChart(), "B2")  # a chart frame with no data

    wb.create_sheet("Notes")["A1"] = "no charts here"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_every_chart_and_picture_once_on_its_own_slide():
    deck, report = convert_workbook_objects(_workbook(), title="Quarterly pack")
    assert report["charts"] == 2  # clustered + stacked; the copy on "Copy" is the same chart
    assert report["pictures"] == 2  # logo once + the orange picture
    assert len(report["duplicates_skipped"]) == 2
    assert report["empty_charts_skipped"] == ["Empty: Chart 1"] or len(report["empty_charts_skipped"]) == 1
    assert report["skipped"] == []

    prs = Presentation(io.BytesIO(deck))
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    pictures = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]
    assert len(charts) == 2 and len(pictures) == 2
    # Native, editable charts: the plotted values and grouping come through unchanged.
    groupings = sorted(c.chart.plots[0]._element.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/chart}grouping").get("val") for c in charts)
    assert groupings == ["clustered", "stacked"]
    assert list(charts[0].chart.plots[0].series[0].values) == [10, 20, 30, 40, 50, 60]
    assert list(charts[0].chart.plots[0].categories) == ["M1", "M2", "M3", "M4", "M5", "M6"]
    headings = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text]
    chart_headings = [h for h in headings if h.startswith("Sales by region")]
    assert chart_headings == ["Sales by region (1 of 2)", "Sales by region (2 of 2)"]
    # Every object fits inside the slide.
    for sh in charts + pictures:
        assert 0 <= sh.left and sh.left + sh.width <= prs.slide_width
        assert 0 <= sh.top and sh.top + sh.height <= prs.slide_height


def test_pictures_keep_original_quality_and_shape():
    deck, _ = convert_workbook_objects(_workbook())
    prs = Presentation(io.BytesIO(deck))
    pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]
    square = [p for p in pics if Image.open(io.BytesIO(p.image.blob)).size == (400, 400)]
    assert square, "the original image file is used, not a re-render"
    assert abs(square[0].width - square[0].height) <= 2  # aspect ratio kept


def test_edit_data_workbook_is_embedded():
    deck, report = convert_workbook_objects(_workbook())
    assert report["edit_data_available"]
    names = zipfile.ZipFile(io.BytesIO(deck)).namelist()
    assert any(n.startswith("ppt/embeddings/") and n.endswith(".xlsx") for n in names)
    assert sum(1 for n in names if n.startswith("ppt/embeddings/")) == 1  # shared, not one copy per chart


def test_keep_duplicates_option():
    _, report = convert_workbook_objects(_workbook(), keep_duplicates=True)
    assert report["charts"] == 3 and report["pictures"] == 3


def test_workbook_without_charts_returns_none():
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    out = io.BytesIO()
    wb.save(out)
    deck, report = convert_workbook_objects(out.getvalue())
    assert deck is None and report["charts"] == 0


os.environ.setdefault("DATABASE_URL", f"sqlite:///{os.path.join(tempfile.gettempdir(), 'meldra_connector_endpoints.db')}")


def test_endpoint_uses_native_charts_and_watermark_keeps_them():
    main = pytest.importorskip("app.main", reason="full backend dependencies not installed")
    from fastapi.testclient import TestClient

    main.init_db()
    main.app.dependency_overrides[main.get_current_user] = lambda: {"email": "ppt-test@example.com", "full_name": "T"}
    try:
        r = TestClient(main.app).post(
            "/api/files/excel-to-ppt",
            files={"file": ("pack.xlsx", _workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    finally:
        main.app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    prs = Presentation(io.BytesIO(r.content))
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart) == 2
