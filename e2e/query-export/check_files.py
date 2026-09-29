"""
Open the downloaded report files the way a customer would and check what is inside:
PDF (PyMuPDF), PowerPoint (python-pptx, then LibreOffice renders it), Word (python-docx,
then LibreOffice), Excel (openpyxl). Writes page images next to the downloads.
Input: one JSON argument {files: {pdf, pptx, docx, xlsx}, charts, titles, out, mode}.
Output: {"checks": [{name, ok, detail}]} on stdout.
"""
import json
import os
import subprocess
import sys
import tempfile

import fitz  # PyMuPDF
import openpyxl
from docx import Document
from pptx import Presentation

args = json.loads(sys.argv[1])
files, charts, titles, out, mode = args["files"], args["charts"], args["titles"], args["out"], args["mode"]
checks = []


def check(name, ok, detail=""):
    checks.append({"name": name, "ok": bool(ok), "detail": detail})


def render(pdf_path, pages, prefix):
    doc = fitz.open(pdf_path)
    for i in pages:
        if i < doc.page_count:
            doc[i].get_pixmap(dpi=70).save(os.path.join(out, f"{mode}-{prefix}-page{i + 1}.png"))
    return doc


def office_to_pdf(path):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, path], check=True, capture_output=True, timeout=180)
    return os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")


# PDF: a cover page, then one page per chart with its picture and data table.
pdf = render(files["pdf"], [0, 1, 2], "pdf")
text = [pdf[i].get_text() for i in range(pdf.page_count)]
check("PDF: cover page plus one page per chart", pdf.page_count == charts + 1, f"{pdf.page_count} pages")
check("PDF: every chart title is in the file", all(any(t in p for p in text) for t in titles))
check("PDF: charts are embedded as pictures", sum(len(pdf[i].get_images()) for i in range(pdf.page_count)) >= charts - 1,
      f"{sum(len(pdf[i].get_images()) for i in range(pdf.page_count))} pictures")

# PowerPoint: title slide, one slide per chart; native (editable) charts where the chart type allows.
prs = Presentation(files["pptx"])
slides = list(prs.slides)
native = sum(1 for s in slides for sh in s.shapes if sh.has_chart)
pictures = sum(1 for s in slides for sh in s.shapes if sh.shape_type == 13)
slide_text = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in slides]
check("PowerPoint: title slide plus one slide per chart", len(slides) == charts + 1, f"{len(slides)} slides")
check("PowerPoint: native, editable charts (not just pictures)", native >= 2, f"{native} native charts, {pictures} pictures")
check("PowerPoint: every chart title on its slide", all(any(t in st for st in slide_text) for t in titles))
check("PowerPoint: opens in LibreOffice", fitz.open(office_to_pdf(files["pptx"])).page_count == len(slides))
render(office_to_pdf(files["pptx"]), [1, 2], "pptx")

# Word: title, then a heading, picture and data table per chart.
doc = Document(files["docx"])
heads = [p.text for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")]
images = len(doc.inline_shapes)
check("Word: a heading and a data table per chart", len(heads) == charts and len(doc.tables) == charts, f"{len(heads)} headings, {len(doc.tables)} tables, {images} pictures")
check("Word: opens in LibreOffice", fitz.open(office_to_pdf(files["docx"])).page_count >= charts)
render(office_to_pdf(files["docx"]), [0, 1], "docx")

# Excel: a contents sheet plus one sheet per chart.
wb = openpyxl.load_workbook(files["xlsx"])
check("Excel: contents sheet plus one sheet per chart", len(wb.sheetnames) == charts + 1, ", ".join(wb.sheetnames))

print(json.dumps({"checks": checks}))
